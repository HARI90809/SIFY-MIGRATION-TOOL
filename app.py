import os
import re
import ipaddress
from io import BytesIO
from pathlib import Path
from functools import wraps
from flask import Flask, render_template, request, redirect, url_for, session, send_file, flash
from openpyxl import load_workbook

app = Flask(__name__)
app.secret_key = os.getenv("FLASK_SECRET", "change-this-secret")
EXCEL_FILE = Path(os.getenv("EXCEL_FILE", str(Path(__file__).parent / "data" / "ESAF IP DETAILS.xlsx")))
BASIC_USER = os.getenv("BASIC_USER", "Sify")
BASIC_PASSWORD = os.getenv("BASIC_PASSWORD", "")
USERS = {BASIC_USER: {"password": BASIC_PASSWORD, "role": "basic"}}

def normalize(value):
    if value is None: return ""
    return " ".join(str(value).strip().split()).lower()

def load_excel_rows():
    if not EXCEL_FILE.exists(): raise FileNotFoundError(f"Excel file not found: {EXCEL_FILE}")
    wb = load_workbook(EXCEL_FILE, data_only=True, read_only=True); ws = wb.active; rows = ws.iter_rows(values_only=True)
    try: headers = next(rows)
    except StopIteration: wb.close(); raise RuntimeError("The Excel workbook is empty.")
    headers = [str(h).strip() if h is not None else "" for h in headers]; records=[]
    for values in rows:
        if not any(v not in (None, "") for v in values): continue
        records.append({h: values[i] if i < len(values) else "" for i,h in enumerate(headers) if h})
    wb.close(); return records

def find_header(headers, *wanted):
    normalized={normalize(h):h for h in headers if h}
    for item in wanted:
        if normalize(item) in normalized: return normalized[normalize(item)]
    def clean(s): return normalize(s).replace("$","").replace("(","").replace(")","")
    cleaned={clean(h):h for h in headers if h}
    for item in wanted:
        if clean(item) in cleaned: return cleaned[clean(item)]
    return None

def get_branch(records, code):
    if not records: return None
    h=find_header(list(records[0].keys()), "BRANCH CODE", "Branch Code", "$(BRANCH CODE)")
    if not h: raise RuntimeError("BRANCH CODE column was not found.")
    target=normalize(code)
    return next((r for r in records if normalize(r.get(h,""))==target), None)

def field(data, *names):
    normalized={normalize(k):k for k in data}
    for name in names:
        key=normalized.get(normalize(name))
        if key is not None and data.get(key) not in (None, ""): return str(data[key]).strip()
    def clean(s): return normalize(s).replace("$","").replace("(","").replace(")","")
    cleaned={clean(k):k for k in data}
    for name in names:
        key=cleaned.get(clean(name))
        if key is not None and data.get(key) not in (None, ""): return str(data[key]).strip()
    return ""

def login_required(fn):
    @wraps(fn)
    def wrapped(*args, **kwargs):
        if "username" not in session: return redirect(url_for("login"))
        return fn(*args, **kwargs)
    return wrapped

def get_sify_mgmt_ip(branch):
    value=field(branch,"$(SIFY_MGMT_IP_SUBNET)","SIFY_MGMT_IP_SUBNET") or field(branch,"$(SIFY_MGMT_IP)","SIFY_MGMT_IP")
    if not value: raise RuntimeError("SIFY-MGMT IP is missing in Excel.")
    return value.split("/",1)[0].strip()

def validate_dhcp_config(switch_mgmt, dhcp_start, dhcp_end):
    if not switch_mgmt or not dhcp_start or not dhcp_end: raise ValueError("FortiLink DHCP values are required in Excel.")
    try:
        gateway=ipaddress.ip_address(str(switch_mgmt).split('/',1)[0]); start=ipaddress.ip_address(str(dhcp_start)); end=ipaddress.ip_address(str(dhcp_end))
    except ValueError as exc: raise ValueError("Invalid FortiLink DHCP values in Excel.") from exc
    subnet=ipaddress.ip_network(f"{gateway}/27",strict=False)
    if start not in subnet or end not in subnet or start>end or start==subnet.network_address or end==subnet.broadcast_address or start<=gateway<=end: raise ValueError(f"Invalid DHCP range {start}-{end} for {subnet}.")
    return str(subnet.netmask)

def normalize_link_status(value):
    return {"BOTH SIFY & MPLS LIVE":"both","BOTH":"both","ONLY SIFY MPLS LIVE":"sify_only","SIFY_ONLY":"sify_only","ONLY AIRTEL ILL LIVE":"airtel_only","AIRTEL_ONLY":"airtel_only"}.get((value or "").strip().upper(),"")

def basic(b, link_status):
    scenario=normalize_link_status(link_status)
    if not scenario: raise ValueError("Please select a valid link status.")
    hostname=field(b,"Hostname"); sdwan=field(b,"SDWAN LINK ID"); sify=field(b,"Sify MP link ID"); bb=field(b,"BB Link ID"); airtel=bb or sdwan
    mgmt=get_sify_mgmt_ip(b); advpn=field(b,"$(BR_LOOPBACK9583)/ADVPN","BR_LOOPBACK9583/ADVPN"); lan=field(b,"$(BR_LAN_INT_IP)"); atm=field(b,"$(ATM_LAN_INTERFACE_IP)"); quarantine=field(b,"$(QUARANTINE_LAN_INT_IP)"); switch_mgmt=field(b,"$(SWITCH_MGMT_LAN_INT_IP)"); dhcp_start=field(b,"$(SWITCH_MGMT_DHCP_START_IP)"); dhcp_end=field(b,"$(SWITCH_MGMT_DHCP_END_IP)"); dhcp_netmask=validate_dhcp_config(switch_mgmt,dhcp_start,dhcp_end); pe_ip=field(b,"$(MPLS1_INT_PE_IP)"); ce_wan_ip=field(b,"$(MPLS1_INT_CE_IP)")
    if scenario in ("both","sify_only"):
        if not ce_wan_ip: raise ValueError("$(MPLS1_INT_CE_IP) is blank or missing in Excel.")
        try: wan_netmask=str(ipaddress.ip_network(f"{ce_wan_ip}/30",strict=False).netmask)
        except ValueError as exc: raise ValueError(f"Invalid $(MPLS1_INT_CE_IP): {ce_wan_ip}") from exc
    lines=[]; add=lines.append
    add('config system global'); add(f'    set hostname "{hostname}"'); add('    set post-login-banner enable'); add('    set pre-login-banner enable'); add('    set switch-controller enable'); add('    set timezone 47'); add('end'); add('')
    add('config system virtual-switch'); add('    edit "lan"'); add('        set physical-switch "sw0"'); add('        config port'); add('            delete "lan3"'); add('        next'); add('    end'); add('end'); add('')
    if scenario in ("both","sify_only"):
        add('config system interface'); add('    edit "wan"'); add('        set vdom "root"'); add('        set mode static'); add(f'        set ip {ce_wan_ip} {wan_netmask}'); add('        set allowaccess ping https ssh'); add('        set type physical'); add(f'        set alias "SIFY-MPLS-{sify}"'); add('        set device-identification enable'); add('        set device-user-identification disable'); add('        set lldp-reception enable'); add('        set lldp-transmission enable'); add('        set monitor-bandwidth enable'); add('        set role wan'); add('    next'); add('end'); add('')
    if scenario in ("both","airtel_only"):
        add('config system interface'); add('    edit "lan3"'); add('        set vdom "root"'); add('        set mode dhcp'); add('        set allowaccess ping https'); add('        set type hard-switch'); add(f'        set alias "AIRTEL-ILL-{airtel}"'); add('        set lldp-reception enable'); add('        set monitor-bandwidth enable'); add('        set role wan'); add('    next'); add('end'); add('')
    add('config system interface'); add('    edit "SIFY-MGMT"'); add('        set vdom "root"'); add(f'        set ip {mgmt} 255.255.255.255'); add('        set allowaccess ping https ssh snmp'); add('        set type loopback'); add('        set alias "SIFY-MGMT-IP"'); add('    next'); add('    edit "loopback9583"'); add('        set vdom "root"'); add(f'        set ip {advpn} 255.255.255.255'); add('        set type loopback'); add('    next'); add('    edit "fortilink"'); add('        set vdom "root"'); add('        set fortilink enable'); add(f'        set ip {switch_mgmt} 255.255.255.224'); add('        set allowaccess ping fabric'); add('        set type aggregate'); add('        set member "a"'); add('        set auto-auth-extension-device enable'); add('        set fortilink-split-interface disable'); add('        set fortilink-neighbor-detect lldp'); add('    next'); add('    edit "_default"'); add('        set vdom "root"'); add(f'        set ip {lan} 255.255.255.128'); add('        set allowaccess ping https ssh'); add('        set alias "LAN"'); add('        set switch-controller-dhcp-snooping enable'); add('        set switch-controller-feature default-vlan'); add('        set interface "fortilink"'); add('        set vlanid 1'); add('    next'); add('    edit "VLAN-20"'); add('        set vdom "root"'); add(f'        set ip {atm} 255.255.255.224'); add('        set allowaccess ping https ssh'); add('        set alias "ATM_LAN"'); add('        set monitor-bandwidth enable'); add('        set role lan'); add('        set interface "fortilink"'); add('        set vlanid 20'); add('    next'); add('    edit "VLAN-30"'); add('        set vdom "root"'); add(f'        set ip {quarantine} 255.255.255.224'); add('        set allowaccess ping https ssh'); add('        set alias "QUARANTINE_VLAN"'); add('        set monitor-bandwidth enable'); add('        set role lan'); add('        set interface "fortilink"'); add('        set vlanid 30'); add('    next'); add('end'); add('')
    add('config system dhcp server'); add('    edit 3'); add('        set dns-service local'); add('        set ntp-service local'); add(f'        set default-gateway {switch_mgmt}'); add(f'        set netmask {dhcp_netmask}'); add('        set interface "fortilink"'); add('        config ip-range'); add('            edit 1'); add(f'                set start-ip {dhcp_start}'); add(f'                set end-ip {dhcp_end}'); add('            next'); add('        end'); add('        set vci-match enable'); add('        set vci-string "FortiSwitch" "FortiExtender"'); add('    next'); add('end'); add('')
    if scenario in ("both","airtel_only"):
        add('config system sdwan'); add('    set status enable'); add('    config zone'); add('        edit "DIA_ZONE"'); add('        next'); add('    end'); add('    config members'); add('        edit 1'); add('            set interface "lan3"'); add('            set zone "DIA_ZONE"'); add('            set gateway 0.0.0.0'); add('        next'); add('    end'); add('end'); add(''); add('config router static'); add('    edit 1'); add('        set sdwan-zone "DIA_ZONE"'); add('    next'); add('end'); add('')
    if scenario=="sify_only":
        if not pe_ip: raise ValueError("$(MPLS1_INT_PE_IP) is required for BGP.")
        remote_as=field(b,"$(BR_MPLS1_REMOTE_AS)","BR_MPLS1_REMOTE_AS") or "9583"
        add('config router bgp'); add('    set as 64665'); add('    config neighbor'); add(f'        edit "{pe_ip}"'); add('            set activate6 disable'); add('            set dont-capability-negotiate enable'); add(f'            set remote-as {remote_as}'); add('            set connect-timer 10'); add('        next'); add('    end'); add('    config network'); add('        edit 2'); add(f'            set prefix {mgmt}/32'); add('        next'); add('    end'); add('end'); add(''); add('config firewall policy'); add('    edit 1'); add('        set srcintf "SIFY-MGMT"'); add('        set dstintf "wan"'); add('        set action accept'); add('        set srcaddr "all"'); add('        set dstaddr "all"'); add('        set schedule "always"'); add('        set service "ALL"'); add('        unset nat'); add('    next'); add('    edit 2'); add('        set srcintf "wan"'); add('        set dstintf "SIFY-MGMT"'); add('        set action accept'); add('        set srcaddr "all"'); add('        set dstaddr "all"'); add('        set schedule "always"'); add('        set service "ALL"'); add('    next'); add('end'); add(''); add('config system central-management'); add('    set type fortimanager'); add('    set fmg "10.240.100.16"'); add(f'    set fmg-source-ip {mgmt}'); add('end'); add(''); add('config system central-management'); add('    set serial-number "FMG-VMTM26002310"'); add('end')
    else:
        add('config system central-management'); add('    set type fortimanager'); add('    set serial-number "FMG-VMTM26002310"'); add('    set fmg "165.101.224.252"'); add('    set interface-select-method specify'); add('    set interface "lan3"'); add('end')
    return "\n".join(lines)+"\n"

@app.route('/login',methods=['GET','POST'])
def login():
    if request.method=='POST':
        u=request.form.get('username',''); p=request.form.get('password',''); user=USERS.get(u)
        if user and user['password'] and p==user['password']: session.clear(); session['username']=u; return redirect(url_for('home'))
        flash('Invalid username or password.','error')
    return render_template('login.html')
@app.get('/logout')
def logout(): session.clear(); return redirect(url_for('login'))
@app.get('/')
@login_required
def home():
    try: status=f"Excel connected · {len(load_excel_rows())} records"
    except Exception as exc: status=str(exc)
    return render_template('index.html',user=session.get('username'),status=status)
@app.post('/lookup')
@login_required
def lookup():
    code=request.form.get('branch_code','').strip()
    try:
        b=get_branch(load_excel_rows(),code)
        if not b: flash(f'Branch code {code} not found in Excel.','error'); return redirect(url_for('home'))
        return render_template('branch.html',branch=b,code=code,user=session.get('username'))
    except Exception as exc: flash(str(exc),'error'); return redirect(url_for('home'))
@app.post('/basic-preview')
@login_required
def basic_preview():
    code=request.form.get('branch_code','').strip(); status=request.form.get('link_status','').strip()
    try:
        b=get_branch(load_excel_rows(),code)
        if not b: return 'Branch not found',404
        return render_template('basic_preview.html',code=code,config=basic(b,status),link_status=status)
    except Exception as exc: return f'Basic config generation failed: {exc}',500
@app.post('/generate')
@login_required
def generate():
    code=request.form.get('branch_code','').strip()
    try:
        b=get_branch(load_excel_rows(),code)
        if not b: return 'Branch not found',404
        host=field(b,'Hostname') or code; text=basic(b,request.form.get('link_status',''))
        return send_file(BytesIO(text.encode()),as_attachment=True,download_name=f'{host}_basic.conf',mimetype='text/plain')
    except Exception as exc: return str(exc),500
@app.get('/health')
def health(): return {'status':'ok','module':'basic-config'}
if __name__=='__main__': app.run(host='0.0.0.0',port=int(os.getenv('PORT','5000')))
